"""Potvin et al. subcortical-volume norms from the distributed workbook.

The workbook stores regression coefficients and covariance matrices in named
ranges. We read those cached numeric values and use formulas in the
``Statistics`` sheet only to associate each displayed region with its model.
No regression values are estimated or substituted by this module.
"""

from __future__ import annotations

import math
import re
from numbers import Real
from typing import Any, Dict, List

import openpyxl
from openpyxl.utils import range_boundaries


_MMULT_RE = re.compile(
    r"\bMMULT\s*\(\s*(X_[A-Za-z0-9_]+)\s*,\s*(B_[A-Za-z0-9_]+)\s*\)",
    flags=re.IGNORECASE,
)


class SubcorticalNorms:
    """Evaluate the original ``mmc2.xlsm`` Potvin normative models."""

    def __init__(self, workbook_path: str):
        self.workbook_path = workbook_path

        # One copy retains formulas/named ranges; the other exposes the numeric
        # values cached in the published workbook.
        self.wb = openpyxl.load_workbook(
            workbook_path, data_only=False, keep_vba=True
        )
        self.wb_vals = openpyxl.load_workbook(
            workbook_path, data_only=True, keep_vba=False
        )

        for sheet_name in ("Statistics", "Matrix"):
            if sheet_name not in self.wb.sheetnames:
                raise ValueError(
                    f"Potvin workbook is missing required sheet {sheet_name!r}"
                )

        self.ws_stats = self.wb["Statistics"]
        self.ws_matrix = self.wb["Matrix"]
        self.ws_matrix_vals = self.wb_vals["Matrix"]

        # openpyxl <= 3.0 exposed ``definedName``; >= 3.1 exposes a mapping.
        defined_names = self.wb.defined_names
        if hasattr(defined_names, "definedName"):
            entries = defined_names.definedName
        else:
            entries = defined_names.values()
        self.defined = {defined_name.name: defined_name for defined_name in entries}

        self.mean_age = self._read_centering_constant("F2", "F6")
        self.mean_tiv = self._read_centering_constant("J2", "J6")
        self.region_models = self._build_region_models()

    # ---------- workbook parsing ----------

    @staticmethod
    def _formula_text(value: Any) -> str:
        """Return formula text for scalar and openpyxl array formulas.

        openpyxl 3.1 represents most formulas in this workbook as
        ``ArrayFormula`` objects, while a few happen to remain strings.
        Treating only strings as formulas silently drops 23 of the 26 models.
        """

        if isinstance(value, str):
            return value
        text = getattr(value, "text", None)
        return text if isinstance(text, str) else ""

    @staticmethod
    def _finite_number(value: Any, *, context: str) -> float:
        if isinstance(value, bool) or not isinstance(value, Real):
            raise ValueError(f"{context} must be numeric, got {value!r}")
        result = float(value)
        if not math.isfinite(result):
            raise ValueError(f"{context} must be finite, got {value!r}")
        return result

    def _read_centering_constant(self, matrix_cell: str, statistics_cell: str) -> float:
        """Read ``Statistics!cell - constant`` from the published formula."""

        formula = self._formula_text(self.ws_matrix[matrix_cell].value)
        pattern = re.compile(
            rf"^\s*=\s*(?:'Statistics'|Statistics)!\$?{statistics_cell[0]}"
            rf"\$?{statistics_cell[1:]}\s*-\s*"
            r"([+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[Ee][+-]?\d+)?)\s*$",
            flags=re.IGNORECASE,
        )
        match = pattern.match(formula)
        if match is None:
            raise ValueError(
                f"Unexpected centering formula in Matrix!{matrix_cell}: {formula!r}"
            )
        return self._finite_number(
            float(match.group(1)), context=f"Matrix!{matrix_cell} centering constant"
        )

    def _defined_range_cells(self, name: str, wb) -> List[List[Any]]:
        """Return a rectangular named range from ``wb`` as values."""

        defined_name = self.defined.get(name)
        if defined_name is None:
            raise ValueError(f"Potvin workbook is missing defined range {name!r}")

        try:
            destinations = list(defined_name.destinations)
        except (AttributeError, TypeError) as exc:
            raise ValueError(f"Defined name {name!r} is not a cell range") from exc
        if len(destinations) != 1:
            raise ValueError(
                f"Defined range {name!r} must have one destination, "
                f"found {len(destinations)}"
            )

        sheet_name, ref = destinations[0]
        if sheet_name not in wb.sheetnames:
            raise ValueError(
                f"Defined range {name!r} refers to missing sheet {sheet_name!r}"
            )
        min_col, min_row, max_col, max_row = range_boundaries(ref)
        ws = wb[sheet_name]
        return [
            [ws.cell(row=row, column=col).value for col in range(min_col, max_col + 1)]
            for row in range(min_row, max_row + 1)
        ]

    def _summary_rows(self) -> Dict[str, Dict[str, float]]:
        """Index the numeric model-summary table at the top of ``Matrix``."""

        rows: Dict[str, Dict[str, float]] = {}
        for row in range(2, self.ws_matrix_vals.max_row + 1):
            key = self.ws_matrix_vals.cell(row=row, column=1).value
            t_value = self.ws_matrix_vals.cell(row=row, column=2).value
            n_value = self.ws_matrix_vals.cell(row=row, column=3).value
            mse_value = self.ws_matrix_vals.cell(row=row, column=4).value
            if not isinstance(key, str):
                continue
            # Model-block rows repeat the key in column A, but have ``X =`` in
            # column B. Only summary rows have numeric t, n, and MSE values.
            if not all(
                isinstance(value, Real) and not isinstance(value, bool)
                for value in (t_value, n_value, mse_value)
            ):
                continue
            if key in rows:
                raise ValueError(f"Duplicate numeric summary row for model {key!r}")
            rows[key] = {
                "t": self._finite_number(t_value, context=f"{key} t value"),
                "n": self._finite_number(n_value, context=f"{key} sample size"),
                "mse": self._finite_number(mse_value, context=f"{key} MSE"),
            }
        return rows

    def _build_region_models(self) -> Dict[str, Dict[str, Any]]:
        """Parse and validate every model displayed in ``Statistics!A11:A36``."""

        models: Dict[str, Dict[str, Any]] = {}
        summaries = self._summary_rows()
        displayed_labels = [
            self.ws_stats.cell(row=row, column=1).value for row in range(11, 37)
        ]
        displayed_labels = [label for label in displayed_labels if label is not None]
        if not displayed_labels:
            raise ValueError("Potvin workbook contains no displayed region labels")

        for row in range(11, 37):
            label = self.ws_stats.cell(row=row, column=1).value
            if label is None:
                continue
            if not isinstance(label, str) or not label.strip():
                raise ValueError(f"Invalid region label in Statistics!A{row}: {label!r}")

            formula = self._formula_text(self.ws_stats.cell(row=row, column=3).value)
            match = _MMULT_RE.search(formula)
            if match is None:
                raise ValueError(
                    f"Could not identify normative MMULT model in Statistics!C{row}: "
                    f"{formula!r}"
                )
            x_name, b_name = match.groups()
            model_key = b_name[2:]
            if x_name[2:].lower() != model_key.lower():
                raise ValueError(
                    f"Statistics!C{row} pairs inconsistent model ranges "
                    f"{x_name!r} and {b_name!r}"
                )
            pred_name = "pred_" + model_key
            matrix_name = "M_" + model_key

            term_rows = self._defined_range_cells(pred_name, self.wb)
            if len(term_rows) != 1:
                raise ValueError(f"{pred_name!r} must be a one-row range")
            beta_names = [value for value in term_rows[0] if value is not None]
            if not beta_names or not all(
                isinstance(value, str) and value.strip() for value in beta_names
            ):
                raise ValueError(f"{pred_name!r} contains invalid covariate names")

            beta_cells = self._defined_range_cells(b_name, self.wb_vals)
            if any(len(values) != 1 for values in beta_cells):
                raise ValueError(f"{b_name!r} must be a one-column range")
            beta_vals = [
                self._finite_number(values[0], context=f"{b_name}[{index}]")
                for index, values in enumerate(beta_cells)
            ]

            matrix_cells = self._defined_range_cells(matrix_name, self.wb_vals)
            matrix_vals = [
                [
                    self._finite_number(
                        value, context=f"{matrix_name}[{row_index},{col_index}]"
                    )
                    for col_index, value in enumerate(values)
                ]
                for row_index, values in enumerate(matrix_cells)
            ]

            dimension = len(beta_names)
            if len(beta_vals) != dimension:
                raise ValueError(
                    f"{label!r} has {dimension} covariates but "
                    f"{len(beta_vals)} coefficients"
                )
            if len(matrix_vals) != dimension or any(
                len(values) != dimension for values in matrix_vals
            ):
                shape = (len(matrix_vals), tuple(len(values) for values in matrix_vals))
                raise ValueError(
                    f"{label!r} covariance matrix has shape {shape}, "
                    f"expected {dimension}x{dimension}"
                )
            for i in range(dimension):
                for j in range(i + 1, dimension):
                    if not math.isclose(
                        matrix_vals[i][j],
                        matrix_vals[j][i],
                        rel_tol=1e-12,
                        abs_tol=1e-15,
                    ):
                        raise ValueError(f"{label!r} covariance matrix is not symmetric")

            summary = summaries.get(model_key)
            if summary is None:
                raise ValueError(f"Missing numeric summary row for model {model_key!r}")
            if summary["n"] <= dimension:
                raise ValueError(f"{label!r} model sample size is not larger than its rank")
            if summary["mse"] <= 0 or summary["t"] <= 0:
                raise ValueError(f"{label!r} has non-positive MSE or t critical value")
            if label in models:
                raise ValueError(f"Duplicate displayed region label {label!r}")

            models[label] = {
                "label": label,
                "model_key": model_key,
                "beta_names": beta_names,
                "beta_vals": beta_vals,
                "M": matrix_vals,
                "n": int(summary["n"]),
                "mse": summary["mse"],
                "t": summary["t"],
                "log10_model": bool(
                    re.search(r"10\s*\^\s*MMULT", formula, flags=re.IGNORECASE)
                ),
            }

        if len(models) != len(displayed_labels):
            raise ValueError(
                f"Parsed {len(models)} of {len(displayed_labels)} displayed Potvin models"
            )
        return models

    # ---------- model evaluation ----------

    def _build_covariates(
        self, age: float, sex: int, field_strength: int, manufacturer: int, icv: float
    ) -> Dict[str, float]:
        """Construct covariates on the scales encoded in ``Matrix!E2:U2``."""

        age = self._finite_number(age, context="age")
        icv = self._finite_number(icv, context="icv")
        if sex not in (0, 1):
            raise ValueError("sex must use the workbook coding 0=female, 1=male")
        if field_strength not in (0, 1):
            raise ValueError("field_strength must use the workbook coding 0=3T, 1=1.5T")
        if manufacturer not in (1, 2, 3):
            raise ValueError("manufacturer must use the workbook coding 1=GE, 2=Philips, 3=Siemens")
        if icv <= 0:
            raise ValueError("icv must be positive")

        agec = age - self.mean_age
        tivc = icv - self.mean_tiv
        sexq = float(sex)
        mfsq = float(field_strength)
        ge = 1.0 if manufacturer == 1 else 0.0
        philips = 1.0 if manufacturer == 2 else 0.0

        return {
            "b0": 1.0,
            "agec": agec,
            "agec2": agec**2,
            "agec3": agec**3,
            "sexq": sexq,
            "tivc": tivc,
            "tivc2": tivc**2,
            "tivc3": tivc**3,
            "mfsq": mfsq,
            "ge": ge,
            "philips": philips,
            "ge_x_mfsq": ge * mfsq,
            "philips_x_mfsq": philips * mfsq,
            "tivc_x_mfsq": tivc * mfsq,
            "agec_x_sexq": agec * sexq,
            "tivc_x_ge": tivc * ge,
            "tivc_x_philips": tivc * philips,
        }

    def regions(self) -> List[str]:
        """Return all region labels displayed by the published workbook."""

        return sorted(self.region_models)

    def predict_region(
        self,
        region_label: str,
        age: float,
        sex: int,
        field_strength: int,
        manufacturer: int,
        icv: float,
    ) -> Dict[str, float]:
        """Compute a normative prediction and its 95% prediction interval."""

        if region_label not in self.region_models:
            raise ValueError(
                f"Unknown region {region_label!r}. Valid regions: {self.regions()}"
            )

        model = self.region_models[region_label]
        covariates = self._build_covariates(
            age, sex, field_strength, manufacturer, icv
        )
        design = []
        for name in model["beta_names"]:
            key = name.lower()
            if key not in covariates:
                raise ValueError(
                    f"Workbook model {region_label!r} uses unsupported covariate {name!r}"
                )
            design.append(covariates[key])

        linear_prediction = sum(
            coefficient * value
            for coefficient, value in zip(model["beta_vals"], design)
        )
        quadratic_form = sum(
            design[i] * model["M"][i][j] * design[j]
            for i in range(len(design))
            for j in range(len(design))
        )
        variance_factor = 1.0 + quadratic_form
        if not math.isfinite(variance_factor) or variance_factor <= 0:
            raise ValueError(
                f"Invalid prediction variance for region {region_label!r}: "
                f"1 + x'Mx = {variance_factor!r}"
            )
        standard_error = math.sqrt(model["mse"] * variance_factor)

        if model["log10_model"]:
            prediction = 10.0**linear_prediction
            lower = 10.0 ** (linear_prediction - model["t"] * standard_error)
            upper = 10.0 ** (linear_prediction + model["t"] * standard_error)
        else:
            prediction = linear_prediction
            lower = linear_prediction - model["t"] * standard_error
            upper = linear_prediction + model["t"] * standard_error

        values = {
            "pred": prediction,
            "lower": lower,
            "upper": upper,
            "se": standard_error,
        }
        for key, value in values.items():
            if not math.isfinite(value):
                raise ValueError(
                    f"Non-finite {key} for region {region_label!r} and supplied covariates"
                )
        return values

    def predict_all(
        self, age: float, sex: int, field_strength: int, manufacturer: int, icv: float
    ) -> Dict[str, Dict[str, float]]:
        """Compute predictions and intervals for all published models."""

        return {
            label: self.predict_region(
                label, age, sex, field_strength, manufacturer, icv
            )
            for label in self.region_models
        }
